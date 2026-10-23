"""Regression tests for document extraction and filename/document search.

Run from the repo root:

    py scripts/test_document_search.py
"""
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import backend.core.hybrid_search as hybrid_search
from backend.core.text_loader import TextLoader
from backend.db.database import Base
from backend.db.models import FileMetadata


def test_text_loader(tmp: Path):
    from docx import Document

    docx_path = tmp / "cv.docx"
    document = Document()
    document.add_heading("Curriculum Vitae", 0)
    document.add_paragraph("Adam IMLOUL")
    document.add_paragraph("Data Scientist")
    document.save(docx_path)

    docx_text = TextLoader.load(str(docx_path))
    assert "Adam IMLOUL" in docx_text, docx_text
    assert "Data Scientist" in docx_text, docx_text

    odt_path = tmp / "cv.odt"
    content_xml = """<?xml version="1.0" encoding="UTF-8"?>
    <office:document-content
        xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
        xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">
      <office:body><office:text>
        <text:p>Curriculum Vitae</text:p>
        <text:p>Data Scientist</text:p>
      </office:text></office:body>
    </office:document-content>"""
    with zipfile.ZipFile(odt_path, "w") as archive:
        archive.writestr("content.xml", content_xml)

    odt_text = TextLoader.load(str(odt_path))
    assert "Curriculum Vitae" in odt_text, odt_text
    assert "Data Scientist" in odt_text, odt_text

    # Legacy .doc is metadata-only: it must never be read as plain text.
    doc_path = tmp / "cv.doc"
    doc_path.write_bytes(b"\xd0\xcf\x11\xe0legacy-binary")
    assert TextLoader.load(str(doc_path)) == ""

    try:
        from reportlab.pdfgen import canvas
    except ImportError:
        print("SKIP: reportlab not installed; PDF extraction not tested")
        return

    pdf_path = tmp / "cv.pdf"
    pdf = canvas.Canvas(str(pdf_path))
    pdf.drawString(72, 720, "Curriculum Vitae - Adam IMLOUL - Data Scientist")
    pdf.save()

    pdf_text = TextLoader.load(str(pdf_path))
    assert "Adam IMLOUL" in pdf_text, pdf_text


def test_filename_document_search():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    TestSession = sessionmaker(autocommit=False, autoflush=False, bind=engine)

    session = TestSession()
    try:
        rows = [
            FileMetadata(
                path=r"C:\Users\me\Desktop\cv_Adam_2026.pdf",
                filename="cv_Adam_2026.pdf",
                extension=".pdf",
                modified_at=datetime(2026, 3, 7),
            ),
            FileMetadata(
                path=r"C:\Users\me\Desktop\resume_Adam.docx",
                filename="resume_Adam.docx",
                extension=".docx",
                modified_at=datetime(2025, 12, 1),
            ),
            FileMetadata(
                path=r"C:\Users\me\Desktop\CV_old.doc",
                filename="CV_old.doc",
                extension=".doc",
                modified_at=datetime(2024, 1, 1),
            ),
            FileMetadata(
                path=r"C:\Users\me\Desktop\curriculum_vitae.pdf",
                filename="curriculum_vitae.pdf",
                extension=".pdf",
                modified_at=datetime(2023, 1, 1),
            ),
            FileMetadata(
                path=r"C:\Users\me\Desktop\helper.txt",
                filename="helper.txt",
                extension=".txt",
                modified_at=datetime(2026, 1, 1),
            ),
            FileMetadata(
                path=r"C:\Users\me\Desktop\course_notes.pdf",
                filename="course_notes.pdf",
                extension=".pdf",
                modified_at=datetime(2026, 1, 1),
            ),
            FileMetadata(
                path=r"C:\Users\me\Desktop\glassfish\cv_noise.pdf",
                filename="cv_noise.pdf",
                extension=".pdf",
                modified_at=datetime(2026, 1, 1),
            ),
        ]
        session.add_all(rows)
        session.commit()
    finally:
        session.close()

    original_session = hybrid_search.SessionLocal
    hybrid_search.SessionLocal = TestSession
    try:
        search = hybrid_search.HybridSearch.__new__(hybrid_search.HybridSearch)

        def paths(query, k=10):
            return [result["path"] for result in search._filename_results(query, k, True)]

        cv_paths = paths("mon cv")
        assert cv_paths[0].endswith("cv_Adam_2026.pdf"), cv_paths
        assert any(path.endswith("resume_Adam.docx") for path in cv_paths), cv_paths
        assert any(path.endswith("CV_old.doc") for path in cv_paths), cv_paths
        assert not any(path.endswith("helper.txt") for path in cv_paths), cv_paths
        assert not any("glassfish" in path for path in cv_paths), cv_paths

        assert paths("mon cv pdf") == [r"C:\Users\me\Desktop\cv_Adam_2026.pdf",
                                        r"C:\Users\me\Desktop\curriculum_vitae.pdf"], paths("mon cv pdf")
        assert paths("mon cv docx") == [r"C:\Users\me\Desktop\resume_Adam.docx"], paths("mon cv docx")
        assert paths("resume")[0].endswith("resume_Adam.docx"), paths("resume")
        assert paths("curriculum vitae")[0].endswith("curriculum_vitae.pdf"), paths("curriculum vitae")
        assert "helper.txt" not in " ".join(paths("pdf")), paths("pdf")
    finally:
        hybrid_search.SessionLocal = original_session
        Base.metadata.drop_all(engine)


def main():
    with tempfile.TemporaryDirectory() as tmp:
        test_text_loader(Path(tmp))
    test_filename_document_search()
    print("PASS: document extraction and filename/document search")


if __name__ == "__main__":
    main()
