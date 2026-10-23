"""Document extraction and filename/document-intent search tests."""

import zipfile
from datetime import datetime

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import backend.core.hybrid_search as hybrid_search
from backend.core.text_loader import TextLoader
from backend.db.database import Base
from backend.db.models import FileMetadata


def test_text_loader_docx_odt_and_legacy_doc(tmp_path):
    from docx import Document

    docx_path = tmp_path / "cv.docx"
    document = Document()
    document.add_heading("Curriculum Vitae", 0)
    document.add_paragraph("Adam IMLOUL")
    document.add_paragraph("Data Scientist")
    document.save(docx_path)

    docx_text = TextLoader.load(str(docx_path))
    assert "Adam IMLOUL" in docx_text
    assert "Data Scientist" in docx_text

    odt_path = tmp_path / "cv.odt"
    content_xml = """<?xml version="1.0" encoding="UTF-8"?>
    <office:document-content
        xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
        xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">
      <office:body><office:text>
        <text:p>Curriculum Vitae</text:p>
        <text:p>Data Scientist</text:p>
      </office:text></office:body>
    </office:document-content>"""
    with zipfile.ZipFile(odt_path, "w") as archive:
        archive.writestr("content.xml", content_xml)

    odt_text = TextLoader.load(str(odt_path))
    assert "Curriculum Vitae" in odt_text
    assert "Data Scientist" in odt_text

    # Legacy .doc is metadata-only and must never be decoded as plain text.
    doc_path = tmp_path / "cv.doc"
    doc_path.write_bytes(b"\xd0\xcf\x11\xe0legacy-binary")
    assert TextLoader.load(str(doc_path)) == ""


def test_text_loader_pdf(tmp_path):
    pytest.importorskip("reportlab.pdfgen")
    from reportlab.pdfgen import canvas

    pdf_path = tmp_path / "cv.pdf"
    pdf = canvas.Canvas(str(pdf_path))
    pdf.drawString(72, 720, "Curriculum Vitae - Adam IMLOUL - Data Scientist")
    pdf.save()

    pdf_text = TextLoader.load(str(pdf_path))
    assert "Adam IMLOUL" in pdf_text


def test_filename_document_search(monkeypatch):
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    test_session = sessionmaker(autocommit=False, autoflush=False, bind=engine)

    session = test_session()
    try:
        session.add_all(
            [
                FileMetadata(
                    path=r"C:\Users\me\Desktop\cv_Adam_2026.pdf",
                    filename="cv_Adam_2026.pdf",
                    extension=".pdf",
                    modified_at=datetime(2026, 3, 7),
                ),
                FileMetadata(
                    path=r"C:\Users\me\Desktop\resume_Adam.docx",
                    filename="resume_Adam.docx",
                    extension=".docx",
                    modified_at=datetime(2025, 12, 1),
                ),
                FileMetadata(
                    path=r"C:\Users\me\Desktop\CV_old.doc",
                    filename="CV_old.doc",
                    extension=".doc",
                    modified_at=datetime(2024, 1, 1),
                ),
                FileMetadata(
                    path=r"C:\Users\me\Desktop\curriculum_vitae.pdf",
                    filename="curriculum_vitae.pdf",
                    extension=".pdf",
                    modified_at=datetime(2023, 1, 1),
                ),
                FileMetadata(
                    path=r"C:\Users\me\Desktop\helper.txt",
                    filename="helper.txt",
                    extension=".txt",
                    modified_at=datetime(2026, 1, 1),
                ),
                FileMetadata(
                    path=r"C:\Users\me\Desktop\course_notes.pdf",
                    filename="course_notes.pdf",
                    extension=".pdf",
                    modified_at=datetime(2026, 1, 1),
                ),
                FileMetadata(
                    path=r"C:\Users\me\Desktop\synthese_grace.docx",
                    filename="synthese_grace.docx",
                    extension=".docx",
                    modified_at=datetime(2026, 1, 1),
                ),
                FileMetadata(
                    path=r"C:\Users\me\Desktop\glassfish\cv_noise.pdf",
                    filename="cv_noise.pdf",
                    extension=".pdf",
                    modified_at=datetime(2026, 1, 1),
                ),
            ]
        )
        session.commit()
    finally:
        session.close()

    monkeypatch.setattr(hybrid_search, "SessionLocal", test_session)
    try:
        search = hybrid_search.HybridSearch.__new__(hybrid_search.HybridSearch)

        def paths(query):
            return [
                result["path"]
                for result in search._filename_results(query, 10, True)
            ]

        cv_paths = paths("mon cv")
        assert cv_paths[0].endswith("cv_Adam_2026.pdf")
        assert any(path.endswith("resume_Adam.docx") for path in cv_paths)
        assert any(path.endswith("CV_old.doc") for path in cv_paths)
        assert not any(path.endswith("helper.txt") for path in cv_paths)
        assert not any(path.endswith("synthese_grace.docx") for path in cv_paths)
        assert not any("glassfish" in path for path in cv_paths)

        assert paths("mon cv pdf") == [
            r"C:\Users\me\Desktop\cv_Adam_2026.pdf",
            r"C:\Users\me\Desktop\curriculum_vitae.pdf",
        ]
        assert paths("mon cv docx") == [
            r"C:\Users\me\Desktop\resume_Adam.docx",
        ]

        resume_paths = paths("resume")
        assert resume_paths[0].endswith("resume_Adam.docx")
        assert paths("curriculum vitae")[0].endswith("curriculum_vitae.pdf")

        pdf_paths = paths("pdf")
        assert any(path.endswith("course_notes.pdf") for path in pdf_paths)
        assert not any(path.endswith("resume_Adam.docx") for path in pdf_paths)
        assert not any(path.endswith("helper.txt") for path in pdf_paths)
    finally:
        Base.metadata.drop_all(engine)
        engine.dispose()
